"""Tests for JSON and Excel export."""

import io
import json

from openpyxl import load_workbook

from app.export import FailedDocument, to_excel, to_json


def test_to_json_has_status_and_values(result):
    data = json.loads(to_json([result, FailedDocument("bad.pdf", "Could not open PDF")]))
    assert data[0]["status"] == "ok"
    assert data[0]["invoice"]["total_amount"] == 12.0
    assert data[0]["invoice"]["invoice_date"] == "2018-12-25"
    assert data[0]["fields"]["total_amount"]["source"] == "ocr_fallback"
    assert data[1] == {"status": "error", "file_name": "bad.pdf", "error": "Could not open PDF"}


def _workbook(results):
    return load_workbook(io.BytesIO(to_excel(results)))


def test_excel_has_expected_sheets_and_rows(result):
    book = _workbook([result, FailedDocument("bad.pdf", "Could not open PDF")])
    assert book.sheetnames == ["Invoices", "Line Items", "Legend"]
    invoices, items = book["Invoices"], book["Line Items"]
    assert invoices.max_row == 3  # header + 2 documents
    assert items.max_row == 3  # header + 2 line items (failed file has none)
    header = [c.value for c in invoices[1]]
    assert header[:4] == ["File", "Status", "Vendor", "Vendor address"]
    row = {h: c.value for h, c in zip(header, invoices[2])}
    assert row["Total"] == 12.0
    assert row["Needs review"] == "Yes"
    assert "total_amount" in row["Fields to check"] and "tax" in row["Fields to check"]
    error_row = {h: c.value for h, c in zip(header, invoices[3])}
    assert error_row["Status"] == "Error" and "Could not open PDF" in error_row["Notes"]


def test_excel_colors_uncertain_cells(result):
    invoices = _workbook([result])["Invoices"]
    header = [c.value for c in invoices[1]]
    total_cell = invoices.cell(row=2, column=header.index("Total") + 1)
    tax_cell = invoices.cell(row=2, column=header.index("Tax") + 1)
    vendor_cell = invoices.cell(row=2, column=header.index("Vendor") + 1)
    assert total_cell.fill.fgColor.rgb.endswith("FFF2CC")  # medium -> yellow
    assert tax_cell.fill.fgColor.rgb.endswith("F8CBAD")  # low -> red
    assert vendor_cell.fill.fill_type is None  # high -> no color
    assert total_cell.number_format == "#,##0.00"


def test_excel_with_no_results_still_valid():
    book = _workbook([])
    assert book["Invoices"].max_row == 1  # header only
