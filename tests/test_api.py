"""API tests with a fake pipeline (no GPU needed)."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.main import create_app, safe_file_name
from tests.conftest import FakePipeline, file_part


def test_health_ok(client):
    response = client.get("/health")
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok" and body["models_loaded"] is True
    assert body["prompt_version"] == "v3"


def test_health_reports_load_error(test_settings):
    with TestClient(create_app(pipeline=FakePipeline(fail_load=True), config=test_settings)) as client:
        response = client.get("/health")
        assert response.status_code == 503
        assert "CUDA not available" in response.json()["error"]
        extract = client.post("/extract", files=[file_part("a.png", b"\x89PNG\r\n\x1a\n...")])
        assert extract.status_code == 503


def test_docs_page(client):
    assert client.get("/docs").status_code == 200
    assert client.get("/", follow_redirects=False).status_code in (302, 307)


def test_extract_png(client, png, fake_pipeline):
    response = client.post("/extract", files=[file_part("receipt.png", png)])
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ok"
    assert body["file_name"] == "receipt.png"
    assert body["invoice"]["total_amount"] == 12.0
    assert body["fields"]["total_amount"]["confidence"] == "medium"
    assert fake_pipeline.processed == ["receipt.png"]


def test_extract_jpg_and_pdf(client, jpg, pdf):
    assert client.post("/extract", files=[file_part("a.jpg", jpg)]).status_code == 200
    assert client.post("/extract", files=[file_part("a.pdf", pdf)]).status_code == 200


def test_rejects_wrong_extension(client):
    response = client.post("/extract", files=[file_part("notes.txt", b"hello")])
    assert response.status_code == 415
    assert response.json()["detail"]["error"] == "unsupported_file_type"


def test_rejects_renamed_file(client):
    response = client.post("/extract", files=[file_part("fake.pdf", b"this is not a pdf")])
    assert response.status_code == 415
    assert response.json()["detail"]["error"] == "file_type_mismatch"


def test_rejects_empty_and_large_files(client):
    assert client.post("/extract", files=[file_part("a.png", b"")]).json()["detail"]["error"] == "empty_file"
    too_big = b"\x89PNG\r\n\x1a\n" + b"0" * (1024 * 1024 + 10)  # limit is 1 MB in test settings
    response = client.post("/extract", files=[file_part("big.png", too_big)])
    assert response.status_code == 413


def test_extraction_error_is_clear(client, png):
    response = client.post("/extract", files=[file_part("broken.png", png)])
    assert response.status_code == 422
    assert response.json()["detail"]["error"] == "extraction_failed"


def test_batch_keeps_going_after_a_bad_file(client, png, pdf):
    files = [file_part("a.png", png, "files"), file_part("broken.png", png, "files"), file_part("c.pdf", pdf, "files")]
    body = client.post("/extract/batch", files=files).json()
    assert [item["status"] for item in body] == ["ok", "error", "ok"]
    assert body[1]["file_name"] == "broken.png" and "valid JSON" in body[1]["error"]


def test_batch_limits(client, png):
    files = [file_part(f"{i}.png", png, "files") for i in range(4)]  # limit is 3 in test settings
    assert client.post("/extract/batch", files=files).status_code == 413


def test_export_excel(client, png, pdf):
    files = [file_part("a.png", png, "files"), file_part("notes.txt", b"x", "files"), file_part("c.pdf", pdf, "files")]
    response = client.post("/export/excel", files=files)
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Invoices", "Line Items", "Legend"]
    statuses = [row[1].value for row in book["Invoices"].iter_rows(min_row=2)]
    assert statuses == ["OK", "Error", "OK"]


def test_safe_file_name():
    assert safe_file_name("../../etc/passwd") == "passwd"
    assert safe_file_name("my receipt (1).jpg") == "my receipt _1_.jpg"
    assert safe_file_name(None) == "upload"
