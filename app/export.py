"""Export results as JSON or as a formatted Excel file.

Excel layout:
    "Invoices"   - one row per document
    "Line Items" - one row per line item, with the file name
    "Legend"     - what the cell colors mean

Values the pipeline is not sure about are colored, so a reviewer can see
at once what to check: yellow = medium confidence (weak match, or taken from
the OCR text by the fallback), red = low confidence (not found on the document).
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from typing import Any, Optional, Sequence, Union

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schema import FIELD_LABELS, LINE_ITEM_LABELS, DocumentResult, FieldResult


@dataclass(frozen=True)
class FailedDocument:
    """A file that could not be processed (kept in the export, so nothing gets lost)."""

    file_name: str
    error: str


ExportItem = Union[DocumentResult, FailedDocument]

# (Excel column name, field name in DocumentResult.fields)
INVOICE_COLUMNS: tuple[tuple[str, str], ...] = tuple(
    (label, name) for name, label in FIELD_LABELS.items()
)
LINE_ITEM_COLUMNS: tuple[tuple[str, str], ...] = tuple(
    (label, name) for name, label in LINE_ITEM_LABELS.items()
)
MONEY_COLUMNS = frozenset({"Subtotal", "Tax", "Total", "Unit price", "Line total"})

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CONFIDENCE_FILLS: dict[str, PatternFill] = {
    "medium": PatternFill("solid", fgColor="FFF2CC"),  # yellow
    "low": PatternFill("solid", fgColor="F8CBAD"),  # red
}
_ERROR_FILL = PatternFill("solid", fgColor="F8CBAD")


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------


def to_dict(item: ExportItem) -> dict[str, Any]:
    """One result as a JSON-friendly dict with a "status" key ("ok" or "error")."""
    if isinstance(item, FailedDocument):
        return {"status": "error", "file_name": item.file_name, "error": item.error}
    return {"status": "ok", **item.model_dump(mode="json")}


def to_json(results: Sequence[ExportItem], indent: Optional[int] = 2) -> str:
    """All results as a JSON string (a list, one entry per document)."""
    return json.dumps([to_dict(item) for item in results], indent=indent, ensure_ascii=False)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def fields_to_check(fields: dict[str, FieldResult]) -> list[str]:
    """Readable names of the fields with medium or low confidence (e.g. "Tax", "Total")."""
    return [
        FIELD_LABELS.get(name, name)
        for name, field in fields.items()
        if field.confidence in ("medium", "low")
    ]


def _invoice_rows(results: Sequence[ExportItem]) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    """Rows for the Invoices sheet, plus the confidence of each cell (for coloring)."""
    rows: list[dict[str, Any]] = []
    cell_confidence: list[dict[str, str]] = []
    for item in results:
        if isinstance(item, FailedDocument):
            row: dict[str, Any] = {"File": item.file_name, "Status": "Error"}
            row.update({column: None for column, _ in INVOICE_COLUMNS})
            row.update({"Needs review": "Yes", "Fields to check": "", "Pages": None,
                        "Processing time (s)": None, "Notes": item.error})
            rows.append(row)
            cell_confidence.append({})
            continue

        row = {"File": item.file_name, "Status": "OK"}
        confidence: dict[str, str] = {}
        for column, field in INVOICE_COLUMNS:
            result = item.fields.get(field)
            row[column] = getattr(item.invoice, field)
            if result is not None and result.confidence:
                confidence[column] = result.confidence
        to_check = fields_to_check(item.fields)
        row.update({
            "Needs review": "Yes" if to_check else "No",
            "Fields to check": ", ".join(to_check),
            "Pages": item.num_pages,
            "Processing time (s)": item.processing_seconds,
            "Notes": " | ".join(item.warnings),
        })
        rows.append(row)
        cell_confidence.append(confidence)
    return rows, cell_confidence


def _line_item_rows(results: Sequence[ExportItem]) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    """Rows for the Line Items sheet, plus the confidence of each cell."""
    rows: list[dict[str, Any]] = []
    cell_confidence: list[dict[str, str]] = []
    for item in results:
        if isinstance(item, FailedDocument):
            continue
        for index, line in enumerate(item.invoice.line_items, start=1):
            fields = item.line_items[index - 1] if index - 1 < len(item.line_items) else {}
            row: dict[str, Any] = {"File": item.file_name, "Line": index}
            confidence: dict[str, str] = {}
            for column, field in LINE_ITEM_COLUMNS:
                row[column] = getattr(line, field)
                result = fields.get(field)
                if result is not None and result.confidence:
                    confidence[column] = result.confidence
            rows.append(row)
            cell_confidence.append(confidence)
    return rows, cell_confidence


def _style_sheet(sheet: Worksheet, columns: list[str], cell_confidence: list[dict[str, str]]) -> None:
    """Header style, column widths, number formats, frozen header, filter and colors."""
    for cell in sheet[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "B2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions

    for col_index, column in enumerate(columns, start=1):
        letter = get_column_letter(col_index)
        values = [str(c.value) for c in sheet[letter][1:] if c.value is not None]
        longest = max([len(column)] + [len(v) for v in values])
        sheet.column_dimensions[letter].width = min(max(12, longest + 2), 60)
        for cell in sheet[letter][1:]:
            if column in MONEY_COLUMNS:
                cell.number_format = "#,##0.00"
            elif column == "Invoice date":
                cell.number_format = "yyyy-mm-dd"
            if column in ("Vendor address", "Notes", "Description"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_index, confidence in enumerate(cell_confidence, start=2):
        for column, level in confidence.items():
            fill = _CONFIDENCE_FILLS.get(level)
            if fill is not None and column in columns:
                sheet.cell(row=row_index, column=columns.index(column) + 1).fill = fill
        if "Status" in columns and sheet.cell(row=row_index, column=columns.index("Status") + 1).value == "Error":
            sheet.cell(row=row_index, column=columns.index("Status") + 1).fill = _ERROR_FILL


def _write_legend(sheet: Worksheet) -> None:
    """A small sheet that explains the colors."""
    rows = [
        ("Color", "Meaning"),
        ("No color", "High confidence: the value was found on the document."),
        ("Yellow", "Medium confidence: weak match, or value taken from the OCR text by the fallback. Please check."),
        ("Red", "Low confidence: the value was not found on the document. Please check."),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    sheet["A3"].fill = _CONFIDENCE_FILLS["medium"]
    sheet["A4"].fill = _CONFIDENCE_FILLS["low"]
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 100


def to_excel(results: Sequence[ExportItem]) -> bytes:
    """All results as an .xlsx file (returned as bytes, ready to save or send)."""
    invoice_rows, invoice_conf = _invoice_rows(results)
    item_rows, item_conf = _line_item_rows(results)

    invoice_columns = ["File", "Status"] + [c for c, _ in INVOICE_COLUMNS] + [
        "Needs review", "Fields to check", "Pages", "Processing time (s)", "Notes"
    ]
    item_columns = ["File", "Line"] + [c for c, _ in LINE_ITEM_COLUMNS]

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(invoice_rows, columns=invoice_columns).to_excel(writer, sheet_name="Invoices", index=False)
        pd.DataFrame(item_rows, columns=item_columns).to_excel(writer, sheet_name="Line Items", index=False)
        _style_sheet(writer.sheets["Invoices"], invoice_columns, invoice_conf)
        _style_sheet(writer.sheets["Line Items"], item_columns, item_conf)
        _write_legend(writer.book.create_sheet("Legend"))
    return buffer.getvalue()
